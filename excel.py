# excel.py
"""
Отчёты и экспорт: XLSX (openpyxl) и PDF (weasyprint).
"""
from flask import Blueprint, request, jsonify, session, send_file, Response, render_template
from database import Database
from datetime import datetime
from utils import login_required, role_required
from logger import get_logger
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

log = get_logger(__name__)

excel_bp = Blueprint('excel', __name__)
db = Database()


# ============================================================
# СТИЛИ XLSX
# ============================================================
HEADER_FILL = PatternFill('solid', fgColor='4E73DF')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=12)
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)

CELL_ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
CELL_ALIGN_CENTER = Alignment(horizontal='center', vertical='center')

BORDER_THIN = Side(border_style='thin', color='D0D0D0')
BORDER = Border(
    left=BORDER_THIN, right=BORDER_THIN,
    top=BORDER_THIN, bottom=BORDER_THIN,
)

STATUS_FILLS = {
    'Новое': PatternFill('solid', fgColor='D6EAF8'),
    'В работе': PatternFill('solid', fgColor='FDEBD0'),
    'Выполнено': PatternFill('solid', fgColor='D5F5E3'),
    'Отменено': PatternFill('solid', fgColor='FADBD8'),
}
PRIORITY_FILLS = {
    'Высокий': PatternFill('solid', fgColor='F5B7B1'),
    'Средний': PatternFill('solid', fgColor='F9E79F'),
    'Низкий': PatternFill('solid', fgColor='ABEBC6'),
}
ZEBRA_FILL = PatternFill('solid', fgColor='F8F9FC')


def _apply_header(ws, headers):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f'A1:{last_col}1'


def _autosize_columns(ws, max_width=60):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for c in col_cells:
            if c.value is not None:
                v = str(c.value)
                l = max(len(line) for line in v.split('\n'))
                if l > max_len:
                    max_len = l
        ws.column_dimensions[letter].width = min(max_len + 3, max_width)


def _style_data_row(ws, row_num, col_count, zebra=False):
    for col in range(1, col_count + 1):
        c = ws.cell(row=row_num, column=col)
        c.border = BORDER
        c.alignment = CELL_ALIGN_LEFT
        if zebra:
            c.fill = ZEBRA_FILL


# ============================================================
# ФИЛЬТРЫ
# ============================================================

def _build_tasks_query(args):
    """Фильтры для экспорта и печати."""
    work_type = args.get('work_type', '')
    cabinet = args.get('cabinet', '')
    status = args.get('status', '')
    user = args.get('user', '')
    date_from = args.get('date_from', '')
    date_to = args.get('date_to', '')

    query = 'SELECT * FROM tasks WHERE deleted_at IS NULL'
    params = []
    if work_type:
        query += ' AND work_type = ?'
        params.append(work_type)
    if cabinet:
        query += ' AND cabinet LIKE ?'
        params.append(f'%{cabinet}%')
    if status:
        query += ' AND status = ?'
        params.append(status)
    if user:
        query += ' AND (from_user LIKE ? OR executor LIKE ? OR assistant LIKE ?)'
        s = f'%{user}%'
        params.extend([s, s, s])
    if date_from:
        query += ' AND deadline >= ?'
        params.append(date_from)
    if date_to:
        query += ' AND deadline <= ?'
        params.append(date_to)
    query += ' ORDER BY created_date DESC'
    return query, params


# ============================================================
# СВОДНЫЙ ОТЧЁТ
# ============================================================

@excel_bp.route('/api/report/tasks')
@login_required
def get_task_report():
    try:
        status_stats = db.query('''
            SELECT status, COUNT(*) AS count FROM tasks
            WHERE deleted_at IS NULL
            GROUP BY status
        ''')
        executor_stats = db.query('''
            SELECT executor,
                   COUNT(*) AS total,
                   SUM(CASE WHEN status = 'Выполнено' THEN 1 ELSE 0 END) AS completed,
                   SUM(CASE WHEN status IN ('Новое', 'В работе') THEN 1 ELSE 0 END) AS active
            FROM tasks WHERE executor != '' AND deleted_at IS NULL
            GROUP BY executor ORDER BY total DESC
        ''')
        monthly_stats = db.query('''
            SELECT strftime('%Y-%m', created_date) AS month,
                   COUNT(*) AS total,
                   SUM(CASE WHEN status = 'Выполнено' THEN 1 ELSE 0 END) AS completed
            FROM tasks WHERE deleted_at IS NULL
            GROUP BY month ORDER BY month DESC LIMIT 12
        ''')
        return jsonify({
            'status_stats': [dict(s) for s in status_stats],
            'executor_stats': [dict(s) for s in executor_stats],
            'monthly_stats': [dict(s) for s in monthly_stats],
        })
    except Exception as e:
        log.exception('Ошибка получения отчёта')
        return jsonify({'error': f'Ошибка: {str(e)}'}), 500


# ============================================================
# ЭКСПОРТ ЗАЯВОК В XLSX
# ============================================================

@excel_bp.route('/api/report/tasks/excel')
@login_required
def export_tasks_excel():
    try:
        query, params = _build_tasks_query(request.args)
        tasks = db.query(query, params)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Заявки'

        headers = [
            '№', 'Дата создания', 'Срок выполнения', 'От кого', 'Кабинет',
            'Описание работы', 'Тип работы', 'Статус', 'Приоритет',
            'Исполнитель', 'Помощник', 'Дата выполнения',
        ]
        _apply_header(ws, headers)

        for i, task in enumerate(tasks, start=1):
            row_num = i + 1
            ws.cell(row=row_num, column=1, value=i)
            ws.cell(row=row_num, column=2, value=task['created_date'] or '')
            ws.cell(row=row_num, column=3, value=task['deadline'] or '')
            ws.cell(row=row_num, column=4, value=task['from_user'] or '')
            ws.cell(row=row_num, column=5, value=task['cabinet'] or '')
            ws.cell(row=row_num, column=6, value=task['description'] or '')
            ws.cell(row=row_num, column=7, value=task['work_type'] or '')
            ws.cell(row=row_num, column=8, value=task['status'] or '')
            ws.cell(row=row_num, column=9, value=task['priority'] or '')
            ws.cell(row=row_num, column=10, value=task['executor'] or '')
            ws.cell(row=row_num, column=11, value=task['assistant'] or '')
            ws.cell(row=row_num, column=12, value=task['completed_date'] or '')

            _style_data_row(ws, row_num, len(headers), zebra=(i % 2 == 0))

            status_cell = ws.cell(row=row_num, column=8)
            if task['status'] in STATUS_FILLS:
                status_cell.fill = STATUS_FILLS[task['status']]
                status_cell.alignment = CELL_ALIGN_CENTER
                status_cell.font = Font(bold=True)

            priority_cell = ws.cell(row=row_num, column=9)
            if task['priority'] in PRIORITY_FILLS:
                priority_cell.fill = PRIORITY_FILLS[task['priority']]
                priority_cell.alignment = CELL_ALIGN_CENTER

            ws.cell(row=row_num, column=1).alignment = CELL_ALIGN_CENTER

        _autosize_columns(ws)

        # Лист статистики
        ws2 = wb.create_sheet('Статистика')
        stats = db.query('''
            SELECT status, COUNT(*) AS count FROM tasks
            WHERE deleted_at IS NULL GROUP BY status
        ''')
        ws2.cell(row=1, column=1, value='Статус').font = Font(bold=True, size=12)
        ws2.cell(row=1, column=2, value='Количество').font = Font(bold=True, size=12)
        for i, s in enumerate(stats, start=2):
            ws2.cell(row=i, column=1, value=s['status'])
            ws2.cell(row=i, column=2, value=s['count'])
        ws2.column_dimensions['A'].width = 25
        ws2.column_dimensions['B'].width = 15

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        log.info(f'XLSX-экспорт заявок: {len(tasks)} записей')

        return send_file(
            output,
            mimetype=(
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
            as_attachment=True,
            download_name=f'tasks_{timestamp}.xlsx',
        )
    except Exception as e:
        log.exception('Ошибка XLSX-экспорта')
        return jsonify({'error': f'Ошибка экспорта: {str(e)}'}), 500


# ============================================================
# ЭКСПОРТ ЗАЯВОК В PDF
# ============================================================

@excel_bp.route('/api/report/tasks/pdf')
@login_required
def export_tasks_pdf():
    """Экспорт заявок в PDF (A4, альбомная ориентация)."""
    try:
        try:
            from weasyprint import HTML
        except ImportError:
            return jsonify({
                'error': 'weasyprint не установлен. '
                         'pip install weasyprint',
            }), 500

        query, params = _build_tasks_query(request.args)
        tasks = db.query(query, params)

        html = render_template(
            'tasks_pdf.html',
            tasks=tasks,
            generated=datetime.now().strftime('%d.%m.%Y %H:%M'),
            total=len(tasks),
        )

        pdf_bytes = HTML(
            string=html,
            base_url=request.url_root,
        ).write_pdf()

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'tasks_{timestamp}.pdf'

        log.info(f'PDF-экспорт заявок: {len(tasks)} записей')

        return Response(
            pdf_bytes,
            mimetype='application/pdf',
            headers={
                'Content-Disposition': f'attachment; filename="{filename}"',
            },
        )
    except Exception as e:
        log.exception('Ошибка PDF-экспорта')
        return jsonify({'error': f'Ошибка: {str(e)}'}), 500


# ============================================================
# ЭКСПОРТ ПОЛЬЗОВАТЕЛЕЙ
# ============================================================

@excel_bp.route('/api/report/users/excel')
@role_required('Администратор')
def export_users_excel():
    try:
        users = db.query('''
            SELECT id, full_name, login, role, email, phone, department, is_active
            FROM users WHERE deleted_at IS NULL
            ORDER BY full_name
        ''')

        wb = Workbook()
        ws = wb.active
        ws.title = 'Пользователи'

        headers = ['ID', 'ФИО', 'Логин', 'Роль', 'Email', 'Телефон', 'Отдел', 'Активен']
        _apply_header(ws, headers)

        for i, u in enumerate(users, start=1):
            row_num = i + 1
            ws.cell(row=row_num, column=1, value=u['id'])
            ws.cell(row=row_num, column=2, value=u['full_name'])
            ws.cell(row=row_num, column=3, value=u['login'])
            ws.cell(row=row_num, column=4, value=u['role'])
            ws.cell(row=row_num, column=5, value=u['email'] or '')
            ws.cell(row=row_num, column=6, value=u['phone'] or '')
            ws.cell(row=row_num, column=7, value=u['department'] or '')
            ws.cell(row=row_num, column=8, value='Да' if u['is_active'] else 'Нет')
            _style_data_row(ws, row_num, len(headers), zebra=(i % 2 == 0))

        _autosize_columns(ws)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return send_file(
            output,
            mimetype=(
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
            as_attachment=True,
            download_name=f'users_{timestamp}.xlsx',
        )
    except Exception as e:
        log.exception('Ошибка XLSX-экспорта пользователей')
        return jsonify({'error': f'Ошибка: {str(e)}'}), 500


# ============================================================
# ДАННЫЕ ДЛЯ ПЕЧАТИ (HTML)
# ============================================================

@excel_bp.route('/api/report/tasks/print')
@login_required
def get_tasks_for_print():
    try:
        query, params = _build_tasks_query(request.args)
        tasks = db.query(query, params)
        return jsonify({
            'tasks': [dict(t) for t in tasks],
            'total': len(tasks),
            'export_date': datetime.now().strftime('%d.%m.%Y %H:%M'),
        })
    except Exception as e:
        log.exception('Ошибка данных для печати')
        return jsonify({'error': f'Ошибка: {str(e)}'}), 500